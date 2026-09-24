import os
import csv
import json
import time
import queue
import threading
import datetime
from pathlib import Path

import customtkinter as ctk
from tkinter import filedialog, messagebox

from ed25519ph_key_generator import generate_keypair
from ed25519ph_signer import sign_pdf_batch
from ed25519ph_verifier import (
    verify, parse_signature_csv, individual_verification,
    ManualValidator, run_rfc8032_test_vectors, run_pdf_dataset_validation,
    save_confusion_matrix_png
)

try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")


class LabeledEntry(ctk.CTkFrame):

    def __init__(self, master, label_text, default="", button_text=None,
                 button_command=None, label_width=110, **kwargs):
        super().__init__(master, fg_color="transparent")
        self.grid_columnconfigure(1, weight=1)

        self.label = ctk.CTkLabel(self, text=label_text, anchor="w", width=label_width)
        self.label.grid(row=0, column=0, sticky="w", padx=(0, 6), pady=4)

        self.entry = ctk.CTkEntry(self, **kwargs)
        if default:
            self.entry.insert(0, default)
        self.entry.grid(row=0, column=1, sticky="ew", pady=4)

        if button_text:
            self.button = ctk.CTkButton(self, text=button_text, width=70,
                                         command=button_command)
            self.button.grid(row=0, column=2, sticky="e", padx=(6, 0), pady=4)

    def get(self):
        return self.entry.get().strip()

    def set(self, value):
        self.entry.delete(0, "end")
        self.entry.insert(0, str(value))


class TitledBox(ctk.CTkFrame):

    def __init__(self, master, title, **kwargs):
        super().__init__(master, border_width=1, border_color="#9a9a9a",
                          corner_radius=4, **kwargs)
        self.grid_columnconfigure(0, weight=1)

        self.title_label = ctk.CTkLabel(
            self, text=title, font=ctk.CTkFont(size=13, weight="bold"), anchor="w"
        )
        self.title_label.grid(row=0, column=0, sticky="w", padx=10, pady=(8, 4))

        self.body = ctk.CTkFrame(self, fg_color="transparent")
        self.body.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        self.grid_rowconfigure(1, weight=1)
        self.body.grid_columnconfigure(0, weight=1)


class ParamDescriptionPanel(ctk.CTkScrollableFrame):

    def __init__(self, master, entries, **kwargs):
        super().__init__(master, label_text="Keterangan Parameter",
                          label_font=ctk.CTkFont(size=13, weight="bold"), **kwargs)
        self.grid_columnconfigure(0, weight=1)

        for i, (heading, body) in enumerate(entries):
            box = ctk.CTkFrame(self, border_width=1, border_color="#9a9a9a", corner_radius=4)
            box.grid(row=i, column=0, sticky="ew", padx=2, pady=4)
            box.grid_columnconfigure(0, weight=1)

            ctk.CTkLabel(box, text=heading, font=ctk.CTkFont(size=12, weight="bold"),
                         anchor="w", justify="left", wraplength=230).grid(
                row=0, column=0, sticky="w", padx=8, pady=(6, 0))
            ctk.CTkLabel(box, text=body, font=ctk.CTkFont(size=11),
                         anchor="w", justify="left", wraplength=230, text_color="#444444").grid(
                row=1, column=0, sticky="w", padx=8, pady=(2, 8))



class App(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Sistem Validasi Dokumen PDF - Ed25519ph (Batch Verification)")
        self.geometry("1280x760")
        self.minsize(1100, 650)

        self.current_keypair = None  
        self.last_signing_csv = None
        self.last_verification_result = None
        self.last_verification_pubkey = None
        self.last_full_validation_report = None  

        self.log_queue_sign = queue.Queue()
        self.log_queue_verify = queue.Queue()
        self.log_queue_validate = queue.Queue()

        self.tabview = ctk.CTkTabview(self)
        self.tabview.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_sign = self.tabview.add("Penandatanganan Dokumen PDF")
        self.tab_verify = self.tabview.add("Batch Verification")
        self.tab_validate = self.tabview.add("Validasi Program")

        self._build_tab_sign(self.tab_sign)
        self._build_tab_verify(self.tab_verify)
        self._build_tab_validate(self.tab_validate)

        self.after(100, self._poll_log_queues)


    def _build_tab_sign(self, parent):
        parent.grid_columnconfigure(1, weight=2)
        parent.grid_columnconfigure((0, 2), weight=1)
        parent.grid_rowconfigure(0, weight=0)
        parent.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(parent, fg_color="transparent")
        header.grid(row=0, column=0, columnspan=3, sticky="ew", padx=8, pady=(8, 0))
        ctk.CTkLabel(header, text="Penandatanganan Dokumen PDF - Ed25519ph",
                      font=ctk.CTkFont(size=18, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(header, text="Pembangkitan kunci, Pre-hash SHA-512, dan Tanda Tangan Ed25519ph",
                      font=ctk.CTkFont(size=12), text_color="#555555").pack(anchor="w")

        left = ctk.CTkFrame(parent, fg_color="transparent")
        left.grid(row=1, column=0, sticky="nsew", padx=8, pady=8)
        left.grid_columnconfigure(0, weight=1)
        left.grid_rowconfigure(2, weight=1)

        param_box = TitledBox(left, "Parameter Input")
        param_box.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        self.entry_dataset = LabeledEntry(
            param_box.body, "Folder dataset", default="testpdf",
            button_text="Pilih...", button_command=self._browse_dataset_folder)
        self.entry_dataset.grid(row=0, column=0, sticky="ew", pady=2)

        self.entry_output = LabeledEntry(
            param_box.body, "Folder output", default="output_signatures",
            button_text="Pilih...", button_command=self._browse_output_folder)
        self.entry_output.grid(row=1, column=0, sticky="ew", pady=2)

        self.entry_ctx = LabeledEntry(param_box.body, "String konteks", default="")
        self.entry_ctx.grid(row=2, column=0, sticky="ew", pady=2)

        self.entry_temper = LabeledEntry(param_box.body, "Temper rate (%)", default="0")
        self.entry_temper.grid(row=3, column=0, sticky="ew", pady=2)

        self.entry_limit = LabeledEntry(param_box.body, "Limit dokumen (0 = semua)", default="0")
        self.entry_limit.grid(row=4, column=0, sticky="ew", pady=2)

        action_box = TitledBox(left, "Tombol Aksi")
        action_box.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        action_box.body.grid_columnconfigure((0, 1), weight=1)

        self.btn_generate_key = ctk.CTkButton(
            action_box.body, text="Generate Kunci Baru", command=self._on_generate_key)
        self.btn_generate_key.grid(row=0, column=0, sticky="ew", padx=(0, 4), pady=4)

        self.btn_sign = ctk.CTkButton(
            action_box.body, text="Mulai Tandatangani", command=self._on_start_signing)
        self.btn_sign.grid(row=0, column=1, sticky="ew", padx=(4, 0), pady=4)

        output_box = TitledBox(left, "Output Hasil")
        output_box.grid(row=2, column=0, sticky="nsew")
        output_box.body.grid_rowconfigure(0, weight=1)

        self.txt_sign_output = ctk.CTkTextbox(output_box.body, wrap="word")
        self.txt_sign_output.grid(row=0, column=0, sticky="nsew")
        self.txt_sign_output.insert(
            "1.0",
            "Belum ada kunci. Klik 'Generate Kunci Baru' untuk membangkitkan "
            "pasangan kunci Ed25519ph (seed -> SHA-512 -> clamping -> scalar s "
            "-> A = [s]B)."
        )
        self.txt_sign_output.configure(state="disabled")

        mid = TitledBox(parent, "Log Proses")
        mid.grid(row=1, column=1, sticky="nsew", padx=8, pady=8)
        mid.body.grid_rowconfigure(0, weight=1)

        self.txt_sign_log = ctk.CTkTextbox(mid.body, wrap="word", font=ctk.CTkFont(family="Consolas", size=11))
        self.txt_sign_log.grid(row=0, column=0, sticky="nsew")

        descriptions = [
            ("Folder dataset",
             "Lokasi folder berisi berkas-berkas PDF yang akan ditandatangani. "
             "Contoh: dataset/pdf/ atau testpdf"),
            ("Folder output",
             "Lokasi folder untuk menyimpan hasil, yaitu berkas "
             "Signature_records.csv yang berisi R, S, pre-hash, dan "
             "informasi tamper tiap dokumen. Contoh: output_signatures/"),
            ("String konteks",
             "Konteks (context string) untuk domain separation Ed25519ph "
             "(dom2). Boleh dikosongkan, atau isi maksimal 255 byte, "
             "misal: 'skripsi-ed25519ph-2026'"),
            ("Temper rate dan contoh",
             "Persentase dokumen yang sengaja dirusak tanda tangannya "
             "(nilai S ditambah 1 mod l) untuk menguji deteksi kesalahan "
             "pada tahap verifikasi. Contoh: 10 berarti 10% dokumen dirusak."),
            ("Limit dokumen dan contoh",
             "Jumlah maksimum berkas PDF yang diproses dari folder dataset. "
             "Isi 0 untuk memproses seluruh berkas PDF yang ditemukan."),
            ("Simpan Kunci Publik (.hex) otomatis",
             "Saat proses penandatanganan selesai, kunci publik A (64 "
             "karakter hex) otomatis disimpan sebagai public_key.hex di "
             "folder output yang sama dengan Signature_records.csv. "
             "Berkas ini dapat langsung dimuat pada tab Batch Verification "
             "menggunakan tombol 'Pilih...'."),
            ("Alur penggunaan",
             "1) Klik 'Generate Kunci Baru' untuk membangkitkan pasangan "
             "kunci Ed25519ph secara acak (seed -> SHA-512 -> clamping "
             "-> scalar s -> A = [s]B).\n"
             "2) Atur parameter folder dataset, folder output, konteks, "
             "temper rate, dan limit dokumen.\n"
             "3) Klik 'Mulai Tandatangani' untuk menghitung pre-hash SHA-512 "
             "tiap PDF dan menandatanganinya dengan Ed25519ph.\n"
             "4) Hasil disimpan otomatis di folder output: "
             "Signature_records.csv dan public_key.hex, siap dipakai "
             "pada tab Batch Verification."),
        ]
        ParamDescriptionPanel(parent, descriptions, width=260).grid(
            row=1, column=2, sticky="nsew", padx=8, pady=8)

    def _build_tab_verify(self, parent):
        parent.grid_columnconfigure(1, weight=2)
        parent.grid_columnconfigure((0, 2), weight=1)
        parent.grid_rowconfigure(0, weight=0)
        parent.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(parent, fg_color="transparent")
        header.grid(row=0, column=0, columnspan=3, sticky="ew", padx=8, pady=(8, 0))
        ctk.CTkLabel(header, text="Batch Verification - Ed25519ph",
                      font=ctk.CTkFont(size=18, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(header, text="Evaluasi persamaan batch dan fallback verifikasi individual",
                      font=ctk.CTkFont(size=12), text_color="#555555").pack(anchor="w")

        left = ctk.CTkFrame(parent, fg_color="transparent")
        left.grid(row=1, column=0, sticky="nsew", padx=8, pady=8)
        left.grid_columnconfigure(0, weight=1)
        left.grid_rowconfigure(2, weight=1)

        param_box = TitledBox(left, "Parameter Input")
        param_box.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        self.entry_csv = LabeledEntry(
            param_box.body, "Signature records", default="output_signatures/Signature_records.csv",
            button_text="Pilih...", button_command=self._browse_csv_file)
        self.entry_csv.grid(row=0, column=0, sticky="ew", pady=2)

        self.entry_pubkey = LabeledEntry(
            param_box.body, "Kunci publik", default="",
            button_text="Pilih...", button_command=self._on_load_pubkey_file_verify)
        self.entry_pubkey.grid(row=1, column=0, sticky="ew", pady=2)

        self.entry_batchsize = LabeledEntry(param_box.body, "Ukuran batch", default="0")
        self.entry_batchsize.grid(row=2, column=0, sticky="ew", pady=2)

        self.entry_ctx_verify = LabeledEntry(param_box.body, "String konteks", default="")
        self.entry_ctx_verify.grid(row=3, column=0, sticky="ew", pady=2)

        action_box = TitledBox(left, "Tombol Aksi")
        action_box.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        action_box.body.grid_columnconfigure((0, 1), weight=1)

        self.btn_verify = ctk.CTkButton(
            action_box.body, text="Mulai Verifikasi", command=self._on_start_verification)
        self.btn_verify.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(4, 8))

        self.btn_save_txt = ctk.CTkButton(
            action_box.body, text="Simpan TXT", command=self._on_save_txt, state="disabled")
        self.btn_save_txt.grid(row=1, column=0, sticky="ew", padx=(0, 4), pady=4)

        self.btn_save_excel = ctk.CTkButton(
            action_box.body, text="Simpan Excel", command=self._on_save_excel, state="disabled")
        self.btn_save_excel.grid(row=1, column=1, sticky="ew", padx=(4, 0), pady=4)

        output_box = TitledBox(left, "Output Hasil Batch Verification")
        output_box.grid(row=2, column=0, sticky="nsew")
        output_box.body.grid_rowconfigure(0, weight=1)

        self.txt_verify_output = ctk.CTkTextbox(output_box.body, wrap="word")
        self.txt_verify_output.grid(row=0, column=0, sticky="nsew")
        self.txt_verify_output.insert(
            "1.0",
            "Belum ada hasil verifikasi. Pilih berkas Signature_records.csv "
            "lalu klik 'Mulai Verifikasi'."
        )
        self.txt_verify_output.configure(state="disabled")

        mid = TitledBox(parent, "Log Proses")
        mid.grid(row=1, column=1, sticky="nsew", padx=8, pady=8)
        mid.body.grid_rowconfigure(0, weight=1)

        self.txt_verify_log = ctk.CTkTextbox(mid.body, wrap="word", font=ctk.CTkFont(family="Consolas", size=11))
        self.txt_verify_log.grid(row=0, column=0, sticky="nsew")

        descriptions = [
            ("Signature records",
             "Lokasi berkas CSV hasil tahap penandatanganan "
             "(Signature_records.csv) yang berisi kolom filename, prehash, "
             "R, S, Signature, dan tampered. Kunci publik TIDAK disimpan "
             "di CSV ini."),
            ("Kunci publik (wajib) dan contoh",
             "Kunci publik A (64 karakter hex / 32 byte) milik penanda "
             "tangan. WAJIB diisi karena CSV tidak menyimpan kunci publik.\n\n"
             "Dua cara mengisi:\n"
             "• Ketik/tempel langsung (64 karakter hex) pada kolom input.\n"
             "• Klik 'Pilih...' untuk memilih berkas public_key.hex yang "
             "disimpan dari tab Penandatanganan — isi field otomatis terisi."),
            ("Ukuran batch dan contoh",
             "Jumlah dokumen yang diproses dalam satu kelompok verifikasi "
             "MSM (Multi-Scalar Multiplication). Isi 0 untuk memverifikasi "
             "seluruh dokumen dalam satu batch. Contoh: 50 untuk membagi "
             "1000 dokumen menjadi 20 batch."),
            ("String konteks dan contoh",
             "Harus SAMA PERSIS dengan konteks yang digunakan saat "
             "penandatanganan, karena nilai ini memengaruhi domain "
             "separator (dom2) dan hash tantangan h."),
            ("Tombol Simpan TXT atau Excel",
             "Setelah verifikasi selesai, hasil ringkasan dan rincian per "
             "dokumen dapat diekspor ke berkas teks (.txt) atau Excel "
             "(.xlsx) dengan sheet Ringkasan, Per Dokumen, dan Log Proses."),
            ("Alur penggunaan",
             "1) Pilih berkas Signature_records.csv dari hasil tab "
             "Penandatanganan.\n"
             "2) Isi kunci publik A (wajib) milik penanda tangan.\n"
             "3) Atur ukuran batch dan string konteks (harus sama dengan "
             "saat menandatangani).\n"
             "4) Klik 'Mulai Verifikasi'. Sistem akan menguji persamaan "
             "agregat batch; jika gagal, sistem otomatis melakukan "
             "verifikasi individual untuk mengisolasi dokumen yang tidak "
             "valid.\n"
             "5) Simpan hasil ke TXT/Excel jika diperlukan."),
            ("Hasil verifikasi",
             "Ringkasan menampilkan jumlah dokumen valid/tidak valid, "
             "daftar nama berkas yang gagal verifikasi, metode (Batch / "
             "Individual), dan waktu proses dalam milidetik."),
        ]
        ParamDescriptionPanel(parent, descriptions, width=260).grid(
            row=1, column=2, sticky="nsew", padx=8, pady=8)

    def _build_tab_validate(self, parent):
        parent.grid_columnconfigure(1, weight=2)
        parent.grid_columnconfigure((0, 2), weight=1)
        parent.grid_rowconfigure(0, weight=0)
        parent.grid_rowconfigure(1, weight=1)

        header = ctk.CTkFrame(parent, fg_color="transparent")
        header.grid(row=0, column=0, columnspan=3, sticky="ew", padx=8, pady=(8, 0))
        ctk.CTkLabel(header, text="Validasi Program - Tahap 1, 2 & 3",
                      font=ctk.CTkFont(size=18, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(header, text="Atur parameter, klik satu tombol — sistem menjalankan "
                                   "Tahap 1 (tes vektor RFC 8032), Tahap 2 (validasi dataset PDF, "
                                   "confusion matrix), dan Tahap 3 (perbandingan waktu & memori) "
                                   "secara berurutan dalam satu Log Proses.",
                      font=ctk.CTkFont(size=12), text_color="#555555").pack(anchor="w")

        left = ctk.CTkFrame(parent, fg_color="transparent")
        left.grid(row=1, column=0, sticky="nsew", padx=8, pady=8)
        left.grid_columnconfigure(0, weight=1)
        left.grid_rowconfigure(2, weight=1)

        param_box = TitledBox(left, "Parameter Input (Tahap 2 & 3)")
        param_box.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        self.entry_dataset_validate = LabeledEntry(
            param_box.body, "Folder dataset PDF", default="testpdf",
            button_text="Pilih...", button_command=self._browse_dataset_folder_validate)
        self.entry_dataset_validate.grid(row=0, column=0, sticky="ew", pady=2)

        self.entry_output_validate = LabeledEntry(
            param_box.body, "Folder output", default="validasi_output",
            button_text="Pilih...", button_command=self._browse_output_folder_validate)
        self.entry_output_validate.grid(row=1, column=0, sticky="ew", pady=2)

        self.entry_ctx_validate = LabeledEntry(param_box.body, "String konteks", default="")
        self.entry_ctx_validate.grid(row=2, column=0, sticky="ew", pady=2)

        self.entry_temper_validate = LabeledEntry(param_box.body, "Temper rate (%)", default="0")
        self.entry_temper_validate.grid(row=3, column=0, sticky="ew", pady=2)

        self.entry_limit_validate = LabeledEntry(param_box.body, "Limit dokumen (0 = semua)", default="0")
        self.entry_limit_validate.grid(row=4, column=0, sticky="ew", pady=2)

        self.entry_batchsize_validate = LabeledEntry(param_box.body, "Ukuran batch (0 = semua)", default="0")
        self.entry_batchsize_validate.grid(row=5, column=0, sticky="ew", pady=2)

        action_box = TitledBox(left, "Tombol Aksi")
        action_box.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        action_box.body.grid_columnconfigure(0, weight=1)

        self.btn_run_full_validation = ctk.CTkButton(
            action_box.body, text="Jalankan Validasi Lengkap (Tahap 1, 2, 3)",
            command=self._on_run_full_validation)
        self.btn_run_full_validation.grid(row=0, column=0, sticky="ew", pady=4)

        self.btn_save_full_report = ctk.CTkButton(
            action_box.body, text="Simpan Laporan Lengkap (TXT)",
            command=self._on_save_full_report, state="disabled")
        self.btn_save_full_report.grid(row=1, column=0, sticky="ew", pady=4)

        output_box = TitledBox(left, "Output Hasil")
        output_box.grid(row=2, column=0, sticky="nsew")
        output_box.body.grid_rowconfigure(0, weight=1)

        self.txt_validate_output = ctk.CTkTextbox(
            output_box.body, wrap="word", font=ctk.CTkFont(family="Consolas", size=11))
        self.txt_validate_output.grid(row=0, column=0, sticky="nsew")
        self.txt_validate_output.insert(
            "1.0",
            "Belum ada hasil. Atur parameter di atas lalu klik\n"
            "'Jalankan Validasi Lengkap (Tahap 1, 2, 3)'."
        )
        self.txt_validate_output.configure(state="disabled")

        self.txt_rfc_output = ctk.CTkTextbox(output_box.body, height=0)
        self.txt_synth_output = ctk.CTkTextbox(output_box.body, height=0)

        mid = TitledBox(parent, "Log Proses")
        mid.grid(row=1, column=1, sticky="nsew", padx=8, pady=8)
        mid.body.grid_rowconfigure(0, weight=1)

        self.txt_validate_log = ctk.CTkTextbox(mid.body, wrap="word", font=ctk.CTkFont(family="Consolas", size=11))
        self.txt_validate_log.grid(row=0, column=0, sticky="nsew")

        descriptions = [
            ("Alur penggunaan",
             "1) Atur 'Folder dataset PDF' (lokasi berkas PDF yang akan "
             "diuji) dan 'Folder output' (tempat menyimpan CSV hasil "
             "tanda tangan sementara).\n"
             "2) Isi 'String konteks', 'Temper rate', 'Limit dokumen', "
             "dan 'Ukuran batch' sesuai kebutuhan.\n"
             "3) Klik 'Jalankan Validasi Lengkap (Tahap 1, 2, 3)' — "
             "sistem menjalankan ketiga tahap secara berurutan dalam "
             "satu Log Proses.\n"
             "4) Hasil setiap tahap muncul langsung di 'Output Hasil' "
             "dan Log Proses tanpa menghapus log sebelumnya.\n"
             "5) Setelah selesai, klik 'Simpan Laporan Lengkap (TXT)' "
             "untuk mengekspor laporan Tahap 1, 2, 3 beserta seluruh "
             "Log Proses ke satu berkas TXT."),
            ("Tahap 1: Tes Vektor RFC 8032",
             "Otomatis dijalankan tanpa parameter tambahan. Sistem "
             "membangkitkan kunci publik dari SECRET KEY resmi RFC 8032 "
             "§7.3 (TEST abc) dan memastikan hasilnya identik dengan "
             "PUBLIC KEY pada RFC (PASS/FAIL), lalu melakukan tanda tangan "
             "+ verifikasi ulang (round-trip)."),
            ("Tahap 2: Validasi Dataset PDF",
             "Menandatangani seluruh dokumen pada 'Folder dataset PDF' "
             "menggunakan kunci sementara dengan 'Temper rate' yang "
             "ditentukan (persentase dokumen yang sengaja dirusak, "
             "S diubah +1 mod l). Batch verification kemudian dijalankan "
             "dan hasilnya dibandingkan dengan kondisi sebenarnya melalui "
             "confusion matrix (TP/TN/FP/FN). Implementasi dinyatakan "
             "VALID jika FP=0 dan FN=0."),
            ("Tahap 3: Uji Performa",
             "Menggunakan CSV dan kunci publik yang dihasilkan otomatis "
             "oleh Tahap 2 untuk membandingkan WAKTU (time.perf_counter, "
             "ms) dan PEAK MEMORY (tracemalloc, MB) antara Verifikasi "
             "Individual dan Batch Verification. Hasil tabel dapat "
             "disalin langsung ke Bab 4 skripsi."),
            ("Folder dataset PDF",
             "Lokasi folder berisi berkas PDF yang akan digunakan "
             "sebagai dataset uji pada Tahap 2 dan 3. "
             "Contoh: dataset/pdf/"),
            ("Folder output",
             "Lokasi folder untuk menyimpan Signature_records.csv hasil "
             "tanda tangan sementara Tahap 2, yang juga digunakan "
             "langsung oleh Tahap 3. Contoh: validasi_output/"),
            ("String konteks",
             "Konteks (context string) untuk domain separation Ed25519ph "
             "(dom2). Boleh dikosongkan. Nilai ini digunakan sama persis "
             "untuk Tahap 2 dan 3 secara otomatis."),
            ("Temper rate (%)",
             "Persentase dokumen yang sengaja dirusak tanda tangannya "
             "(S diubah +1 mod l) untuk membentuk kelompok 'dirusak' "
             "pada Tahap 2. Contoh: 50 = separuh dokumen dirusak."),
            ("Limit dokumen",
             "Jumlah maksimum berkas PDF yang diproses dari folder "
             "dataset. Isi 0 untuk memproses seluruh berkas PDF."),
            ("Ukuran batch",
             "Jumlah dokumen per kelompok pada Batch Verification (MSM) "
             "di Tahap 2 dan 3. Isi 0 untuk satu batch penuh."),
        ]
        ParamDescriptionPanel(parent, descriptions, width=260).grid(
            row=1, column=2, sticky="nsew", padx=8, pady=8)


    def _browse_dataset_folder(self):
        path = filedialog.askdirectory(title="Pilih folder dataset PDF")
        if path:
            self.entry_dataset.set(path)

    def _browse_output_folder(self):
        path = filedialog.askdirectory(title="Pilih folder output")
        if path:
            self.entry_output.set(path)

    def _browse_csv_file(self):
        path = filedialog.askopenfilename(
            title="Pilih berkas Signature_records.csv",
            filetypes=[("CSV files", "*.csv"), ("Semua berkas", "*.*")])
        if path:
            self.entry_csv.set(path)

    def _browse_dataset_folder_validate(self):
        path = filedialog.askdirectory(title="Pilih folder dataset PDF")
        if path:
            self.entry_dataset_validate.set(path)

    def _browse_output_folder_validate(self):
        path = filedialog.askdirectory(title="Pilih folder output")
        if path:
            self.entry_output_validate.set(path)

    @staticmethod
    def _read_pubkey_hex_file(path: str) -> str:
        """
        Baca berkas .hex dan kembalikan string hex kunci publik (64 karakter).
        Berkas boleh berisi spasi/newline — hanya karakter hex yang diambil.
        Validasi: hasil harus tepat 64 karakter hex (32 byte).
        """
        with open(path, 'r', encoding='utf-8') as f:
            raw = f.read()
        cleaned = ''.join(c for c in raw.lower() if c in '0123456789abcdef')
        if len(cleaned) != 64:
            raise ValueError(
                f"Berkas tidak berisi kunci publik Ed25519 yang valid.\n"
                f"Diharapkan 64 karakter hex (32 byte), ditemukan {len(cleaned)} karakter."
            )
        return cleaned

    def _on_load_pubkey_file_verify(self):
        """Muat kunci publik dari berkas .hex ke field input kunci publik Tab 2."""
        path = filedialog.askopenfilename(
            title="Muat kunci publik dari berkas .hex",
            filetypes=[("Hex files", "*.hex"), ("Text files", "*.txt"), ("Semua berkas", "*.*")]
        )
        if not path:
            return
        try:
            pubkey_hex = self._read_pubkey_hex_file(path)
            self.entry_pubkey.set(pubkey_hex)
            self.log_queue_verify.put(f"[INFO] Kunci publik dimuat dari: {path}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _poll_log_queues(self):
        self._drain_queue(self.log_queue_sign, self.txt_sign_log)
        self._drain_queue(self.log_queue_verify, self.txt_verify_log)
        self._drain_queue(self.log_queue_validate, self.txt_validate_log)
        self.after(100, self._poll_log_queues)

    @staticmethod
    def _drain_queue(q, textbox):
        updated = False
        while True:
            try:
                msg = q.get_nowait()
            except queue.Empty:
                break
            textbox.configure(state="normal")
            textbox.insert("end", msg + "\n")
            updated = True
        if updated:
            textbox.see("end")
            textbox.configure(state="normal")

    def _set_textbox(self, textbox, content):
        textbox.configure(state="normal")
        textbox.delete("1.0", "end")
        textbox.insert("1.0", content)
        textbox.configure(state="disabled")

    def _append_textbox(self, textbox, content):
        textbox.configure(state="normal")
        textbox.insert("end", content)
        textbox.see("end")
        textbox.configure(state="disabled")

    def _on_generate_key(self):
        try:
            self.current_keypair = generate_keypair()
            self._show_keypair_info("dibangkitkan secara acak (os.urandom)")
            self.log_queue_sign.put("[INFO] Kunci baru berhasil dibangkitkan.")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal membangkitkan kunci:\n{e}")

    def _show_keypair_info(self, sumber_text):
        kp = self.current_keypair
        info = (
            f"Pasangan kunci {sumber_text}.\n\n"
            f"Seed (hex)        : {kp['seed_hex']}\n"
            f"Scalar s          : {kp['scalar_s']}\n"
            f"Prefix (hex)      : {kp['prefix']}\n"
            f"Kunci Privat (hex): {kp['privkey_hex']}\n"
            f"Kunci Publik A    : {kp['pubkey_hex']}\n"
        )
        self._set_textbox(self.txt_sign_output, info)

        self.entry_pubkey.set(kp['pubkey_hex'])

    def _on_start_signing(self):
        if self.current_keypair is None:
            messagebox.showwarning(
                "Peringatan",
                "Belum ada kunci. Klik 'Generate Kunci Baru' atau muat dari seed terlebih dahulu.")
            return

        dataset_folder = self.entry_dataset.get()
        output_folder = self.entry_output.get()
        ctx = self.entry_ctx.get()

        try:
            temper_rate = float(self.entry_temper.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Temper rate harus berupa angka.")
            return

        try:
            limit = int(self.entry_limit.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Limit dokumen harus berupa bilangan bulat.")
            return

        if not os.path.isdir(dataset_folder):
            messagebox.showerror("Error", f"Folder dataset tidak ditemukan:\n{dataset_folder}")
            return

        self.btn_sign.configure(state="disabled", text="Memproses...")
        self.txt_sign_log.configure(state="normal")
        self.txt_sign_log.delete("1.0", "end")
        self.txt_sign_log.configure(state="disabled")

        privkey_hex = self.current_keypair['privkey_hex']

        thread = threading.Thread(
            target=self._signing_worker,
            args=(dataset_folder, privkey_hex, ctx, temper_rate, limit, output_folder),
            daemon=True
        )
        thread.start()

    def _signing_worker(self, dataset_folder, privkey_hex, ctx, temper_rate, limit, output_folder):
        def log(msg):
            self.log_queue_sign.put(str(msg))

        try:
            log("=== Memulai proses penandatanganan ===")
            start = time.perf_counter()
            csv_path, peak_mb = sign_pdf_batch(
                pdf_folder=dataset_folder,
                privkey_hex=privkey_hex,
                ctx=ctx,
                temper_rate=temper_rate,
                limit=limit,
                output_folder=output_folder,
                log_callback=log
            )
            elapsed = (time.perf_counter() - start) * 1000
            self.last_signing_csv = csv_path
            log(f"=== Selesai dalam {elapsed:.2f} ms ===")
            log(f"[INFO] Peak Memory saat Pre-hash & Sign: {peak_mb:.4f} MB")

            kp = self.current_keypair
            pubkey_path = str(Path(output_folder) / "public_key.hex")
            with open(pubkey_path, 'w', encoding='utf-8') as f:
                f.write(kp['pubkey_hex'])
            log(f"[INFO] Kunci publik disimpan ke: {pubkey_path}")

            summary = (
                f"Penandatanganan selesai.\n\n"
                f"Kunci Publik A     : {kp['pubkey_hex']}\n"
                f"Folder dataset     : {dataset_folder}\n"
                f"Folder output      : {output_folder}\n"
                f"String konteks     : '{ctx}'\n"
                f"Temper rate        : {temper_rate}%\n"
                f"Limit dokumen      : {limit if limit > 0 else 'semua'}\n"
                f"Berkas CSV         : {csv_path}\n"
                f"Kunci publik (.hex): {pubkey_path}\n"
                f"Waktu proses       : {elapsed:.2f} ms\n"
                f"Peak Memory (RAM)  : {peak_mb:.4f} MB\n"
            )
            self.after(0, lambda: self._set_textbox(self.txt_sign_output, summary))

            self.after(0, lambda: self.entry_csv.set(csv_path))
            self.after(0, lambda: self.entry_dataset_validate.set(dataset_folder))

        except Exception as e:
            log(f"[ERROR] {e}")
            self.after(0, lambda: messagebox.showerror("Error", f"Proses gagal:\n{e}"))
        finally:
            self.after(0, lambda: self.btn_sign.configure(state="normal", text="Mulai Tandatangani"))

    def _on_start_verification(self):
        csv_path = self.entry_csv.get()
        if not os.path.isfile(csv_path):
            messagebox.showerror("Error", f"Berkas CSV tidak ditemukan:\n{csv_path}")
            return

        ctx = self.entry_ctx_verify.get()

        try:
            batch_size = int(self.entry_batchsize.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Ukuran batch harus berupa bilangan bulat.")
            return

        pubkey_hex = self.entry_pubkey.get().strip().lower()
        if not pubkey_hex:
            messagebox.showwarning(
                "Peringatan",
                "Kunci publik (hex) wajib diisi. CSV tidak menyimpan kunci publik; "
                "isi kunci publik A milik penanda tangan (lihat tab Penandatanganan).")
            return
        try:
            pubkey_bytes = bytes.fromhex(pubkey_hex)
            if len(pubkey_bytes) != 32:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Kunci publik harus berupa hex 64 karakter (32 byte).")
            return

        self.btn_verify.configure(state="disabled", text="Memproses...")
        self.btn_save_txt.configure(state="disabled")
        self.btn_save_excel.configure(state="disabled")
        self.txt_verify_log.configure(state="normal")
        self.txt_verify_log.delete("1.0", "end")
        self.txt_verify_log.configure(state="disabled")

        thread = threading.Thread(
            target=self._verification_worker,
            args=(csv_path, ctx, batch_size, pubkey_hex),
            daemon=True
        )
        thread.start()

    def _verification_worker(self, csv_path, ctx, batch_size, pubkey_hex):
        def log(msg):
            self.log_queue_verify.put(str(msg))

        try:
            log("=== Memulai proses verifikasi ===")
            log(f"[INFO] Kunci publik A (diinput dari GUI): {pubkey_hex}")

            result = verify(csv_path, pubkey_hex=pubkey_hex, ctx=ctx, batch_size=batch_size, log_callback=log)
            self.last_verification_result = result
            self.last_verification_ctx = ctx
            self.last_verification_pubkey = pubkey_hex

            valid_pct = (result.valid / result.total * 100) if result.total else 0.0
            summary = (
                f"Hasil Batch Verification\n\n"
                f"Berkas CSV       : {csv_path}\n"
                f"Kunci publik A   : {pubkey_hex}\n"
                f"String konteks   : '{ctx}'\n"
                f"Ukuran batch     : {batch_size if batch_size > 0 else 'semua (1 batch)'}\n"
                f"Metode           : {result.method}\n"
                f"Total dokumen    : {result.total}\n"
                f"Valid            : {result.valid} ({valid_pct:.2f}%)\n"
                f"Tidak valid      : {result.invalid}\n"
                f"Waktu proses     : {result.time_taken:.2f} ms\n"
            )
            if result.invalid_files:
                summary += "\nDokumen tidak valid:\n"
                for fn in result.invalid_files:
                    summary += f"  - {fn}\n"

            self.after(0, lambda: self._set_textbox(self.txt_verify_output, summary))
            self.after(0, lambda: self.btn_save_txt.configure(state="normal"))
            self.after(0, lambda: self.btn_save_excel.configure(
                state="normal" if OPENPYXL_AVAILABLE else "disabled"))

        except Exception as e:
            log(f"[ERROR] {e}")
            self.after(0, lambda: messagebox.showerror("Error", f"Proses gagal:\n{e}"))
        finally:
            self.after(0, lambda: self.btn_verify.configure(state="normal", text="Mulai Verifikasi"))

    def _on_save_txt(self):
        result = self.last_verification_result
        if result is None:
            messagebox.showwarning("Peringatan", "Belum ada hasil verifikasi untuk disimpan.")
            return

        path = filedialog.asksaveasfilename(
            title="Simpan hasil sebagai TXT",
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt")],
            initialfile="Hasil_Verifikasi.txt"
        )
        if not path:
            return

        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write("LAPORAN BATCH VERIFICATION - Ed25519ph\n")
                f.write(f"Tanggal           : {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"Berkas CSV        : {self.entry_csv.get()}\n")
                f.write(f"Kunci publik A    : {getattr(self, 'last_verification_pubkey', '')}\n")
                f.write(f"String konteks    : '{getattr(self, 'last_verification_ctx', '')}'\n")
                f.write(f"Metode            : {result.method}\n")
                f.write(f"Total dokumen     : {result.total}\n")
                f.write(f"Valid             : {result.valid}\n")
                f.write(f"Tidak valid       : {result.invalid}\n")
                f.write(f"Waktu proses (ms) : {result.time_taken:.4f}\n\n")

                f.write("=== RINCIAN PER DOKUMEN ===\n")
                invalid_set = set(result.invalid_files)
                for rec in result.records:
                    status = "TIDAK VALID" if rec.filename in invalid_set else "VALID"
                    f.write(f"{rec.filename:40s} : {status}\n")

                f.write("\n=== LOG PROSES ===\n")
                f.write(self.txt_verify_log.get("1.0", "end"))

            messagebox.showinfo("Berhasil", f"Hasil disimpan ke:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal menyimpan TXT:\n{e}")

    def _on_save_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Error", "Modul openpyxl tidak tersedia.")
            return

        result = self.last_verification_result
        if result is None:
            messagebox.showwarning("Peringatan", "Belum ada hasil verifikasi untuk disimpan.")
            return

        path = filedialog.asksaveasfilename(
            title="Simpan hasil sebagai Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="Hasil_Verifikasi.xlsx"
        )
        if not path:
            return

        try:
            wb = Workbook()

            ws_summary = wb.active
            ws_summary.title = "Ringkasan"
            valid_pct = (result.valid / result.total * 100) if result.total else 0.0
            rows = [
                ("Tanggal", datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
                ("Berkas CSV", self.entry_csv.get()),
                ("Kunci publik A", getattr(self, 'last_verification_pubkey', '')),
                ("String konteks", getattr(self, 'last_verification_ctx', '')),
                ("Metode", result.method),
                ("Total dokumen", result.total),
                ("Valid", result.valid),
                ("Tidak valid", result.invalid),
                ("Persentase valid (%)", round(valid_pct, 4)),
                ("Waktu proses (ms)", round(result.time_taken, 4)),
            ]
            for r in rows:
                ws_summary.append(r)

            ws_docs = wb.create_sheet("Per Dokumen")
            ws_docs.append(["No", "Nama Berkas", "Status", "Ditandai Tampered (saat sign)"])
            invalid_set = set(result.invalid_files)
            for i, rec in enumerate(result.records, start=1):
                status = "TIDAK VALID" if rec.filename in invalid_set else "VALID"
                ws_docs.append([i, rec.filename, status, "Ya" if rec.tampered else "Tidak"])

            ws_log = wb.create_sheet("Log Proses")
            ws_log.append(["Log"])
            for line in self.txt_verify_log.get("1.0", "end").splitlines():
                if line.strip():
                    ws_log.append([line])

            wb.save(path)
            messagebox.showinfo("Berhasil", f"Hasil disimpan ke:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal menyimpan Excel:\n{e}")

    def _on_run_full_validation(self):
        pdf_folder = self.entry_dataset_validate.get()
        if not os.path.isdir(pdf_folder):
            messagebox.showerror("Error", f"Folder dataset tidak ditemukan:\n{pdf_folder}")
            return

        output_folder = self.entry_output_validate.get()
        ctx = self.entry_ctx_validate.get()

        try:
            temper_rate = float(self.entry_temper_validate.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Temper rate harus berupa angka.")
            return

        try:
            limit = int(self.entry_limit_validate.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Limit dokumen harus berupa bilangan bulat.")
            return

        try:
            batch_size = int(self.entry_batchsize_validate.get() or "0")
        except ValueError:
            messagebox.showerror("Error", "Ukuran batch harus berupa bilangan bulat.")
            return

        self.btn_run_full_validation.configure(state="disabled", text="Memproses...")
        self.btn_save_full_report.configure(state="disabled")

        self.txt_validate_log.configure(state="normal")
        self.txt_validate_log.delete("1.0", "end")
        self.txt_validate_log.configure(state="disabled")
        self._set_textbox(self.txt_validate_output, "")

        self._set_textbox(self.txt_rfc_output, "Memproses Tahap 1...")
        self._set_textbox(self.txt_synth_output, "Memproses Tahap 2...")
        self._set_textbox(self.txt_validate_output, "Menunggu Tahap 2 selesai...")

        thread = threading.Thread(
            target=self._full_validation_worker,
            args=(pdf_folder, output_folder, ctx, temper_rate, limit, batch_size),
            daemon=True
        )
        thread.start()

    def _full_validation_worker(self, pdf_folder, output_folder, ctx, temper_rate, limit, batch_size):
        def log(msg):
            self.log_queue_validate.put(str(msg))

        try:
            log("##### TAHAP 1: TES VEKTOR RFC 8032 (Ed25519ph) #####")
            vector_results = run_rfc8032_test_vectors(log_callback=log)

            lines1 = []
            for r in vector_results:
                status = "PASS" if r.passed else "FAIL"
                lines1.append(f"[{status}] {r.name}")
                for line in r.detail.splitlines():
                    lines1.append(f"    {line}")
                lines1.append("")
            tahap1_passed = all(r.passed for r in vector_results)
            lines1.append(
                f"Kesimpulan: {'SEMUA VEKTOR LOLOS (PASS)' if tahap1_passed else 'ADA VEKTOR YANG GAGAL (FAIL)'}"
            )
            tahap1_report = "\n".join(lines1)
            self.after(0, lambda r=tahap1_report: self._append_textbox(
                self.txt_validate_output, "=== TAHAP 1: TES VEKTOR RFC 8032 ===\n" + r))
            log("")

            log("##### TAHAP 2: VALIDASI DATASET PDF (VALID vs DIRUSAK) #####")
            tahap2_result = run_pdf_dataset_validation(
                pdf_folder=pdf_folder,
                ctx=ctx,
                temper_rate=temper_rate,
                limit=limit,
                batch_size=batch_size,
                output_folder=output_folder,
                log_callback=log
            )
            self.after(0, lambda r=tahap2_result.report: self._append_textbox(
                self.txt_validate_output, "\n\n=== TAHAP 2: VALIDASI DATASET PDF ===\n" + r))

            cm_path = str(Path(output_folder) / "confusion_matrix.png")
            save_confusion_matrix_png(tahap2_result, cm_path, log_callback=log)
            log("")

            log("##### TAHAP 3: UJI PERFORMA (WAKTU & MEMORI) #####")
            log(f"Berkas CSV   : {tahap2_result.csv_path}")
            log(f"Kunci publik : {tahap2_result.pubkey_hex}")
            log(f"Konteks      : '{ctx}'")
            log(f"Ukuran batch : {batch_size if batch_size > 0 else 'semua (1 batch)'}")
            log("")

            validator = ManualValidator(
                tahap2_result.csv_path,
                pubkey_hex=tahap2_result.pubkey_hex,
                ctx=ctx,
                pdf_folder=pdf_folder
            )

            log("[Tahap 3 - Fase 1] Mengukur peak RAM per dokumen saat pre-hash & signing...")
            fase1_rows = validator.run_signing_memory_phase(log_callback=log)

            if fase1_rows:
                fase1_lines = [
                    "--- FASE 1: PENANDATANGANAN (Peak RAM per Dokumen, Pre-hash SHA-512) ---",
                    f"Jumlah dokumen diukur : {len(fase1_rows)}",
                    f"Catatan               : Peak RAM diukur per dokumen (tracemalloc reset per iterasi).",
                    f"                        Membuktikan RAM stabil meskipun ukuran PDF bervariasi.",
                    "",
                    f"{'No':<5}{'Nama File':<35}{'Ukuran (KB)':<14}{'Peak RAM (MB)':<16}{'Waktu (ms)':<12}",
                ]
                for i, row in enumerate(fase1_rows, 1):
                    fase1_lines.append(
                        f"{i:<5}{row['filename']:<35}{row['size_kb']:<14.2f}"
                        f"{row['peak_mb']:<16.6f}{row['time_ms']:<12.4f}"
                    )
                peak_vals = [r['peak_mb'] for r in fase1_rows]
                size_vals = [r['size_kb'] for r in fase1_rows]
                fase1_lines += [
                    "",
                    f"Ukuran file (min/max) : {min(size_vals):.2f} KB / {max(size_vals):.2f} KB",
                    f"Peak RAM (min/max/avg): {min(peak_vals):.6f} MB / {max(peak_vals):.6f} MB "
                    f"/ {sum(peak_vals)/len(peak_vals):.6f} MB",
                    f"Variasi peak RAM      : {max(peak_vals) - min(peak_vals):.6f} MB "
                    f"(semakin kecil = semakin stabil)",
                ]
                fase1_report = "\n".join(fase1_lines)
                for line in fase1_lines:
                    log(line)
            else:
                fase1_report = "[Fase 1] Tidak ada berkas PDF ditemukan di folder dataset untuk diukur."
                log(fase1_report)

            log("")

            log("[Tahap 3 - Fase 2] Mengukur waktu & peak RAM verifikasi (individual vs batch)...")
            tahap3_comparison = validator.compare(batch_size=batch_size, log_callback=log)
            log("")

            tahap3_full = fase1_report + "\n\n" + tahap3_comparison['report']

            self.after(0, lambda r=tahap3_full: self._append_textbox(
                self.txt_validate_output, "\n\n=== TAHAP 3: UJI PERFORMA (WAKTU & MEMORI) ===\n" + r))

            full_report = "\n\n".join([tahap1_report, tahap2_result.report, tahap3_full])
            self.last_full_validation_report = full_report
            log("##### VALIDASI LENGKAP (TAHAP 1, 2, 3) SELESAI #####")

            self.after(0, lambda: self.btn_save_full_report.configure(state="normal"))

        except Exception as e:
            log(f"[ERROR] {e}")
            self.after(0, lambda: messagebox.showerror("Error", f"Proses gagal:\n{e}"))
        finally:
            self.after(0, lambda: self.btn_run_full_validation.configure(
                state="normal", text="Jalankan Validasi Lengkap (Tahap 1, 2, 3)"))

    def _on_save_full_report(self):
        report = self.last_full_validation_report
        if not report:
            messagebox.showwarning("Peringatan", "Belum ada hasil validasi untuk disimpan.")
            return

        path = filedialog.asksaveasfilename(
            title="Simpan laporan validasi lengkap sebagai TXT",
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt")],
            initialfile="Laporan_Validasi_Skripsi.txt"
        )
        if not path:
            return

        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write("LAPORAN VALIDASI PROGRAM LENGKAP (TAHAP 1, 2, 3) - Ed25519ph\n")
                f.write(f"Tanggal: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
                f.write(report)

            messagebox.showinfo("Berhasil", f"Laporan disimpan ke:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Gagal menyimpan laporan:\n{e}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
